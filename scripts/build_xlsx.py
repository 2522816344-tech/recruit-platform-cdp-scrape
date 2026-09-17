# -*- coding: utf-8 -*-
"""合并 猎聘 + 前程无忧 + 智联招聘 + BOSS直聘 的抓取结果 → 多 Sheet Excel

工作目录优先取环境变量 SCRAPE_DIR，否则取进程当前目录。
标题 / 本地城市 / 剔除规则全部从工作目录的 config.json 读，缺省时用内置默认值。
"""
import json, re, os, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Windows 编码兜底：英文版 Windows 控制台是 cp1252，print 中文/✓✗ 会 UnicodeEncodeError。
# 统一切到 UTF-8 输出（Python 3.7+）；errors='replace' 保证极端情况不崩。
# ---------------------------------------------------------------------------
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

BASE = os.environ.get('SCRAPE_DIR') or os.getcwd()
DATA = os.environ.get('SCRAPE_DATA') or os.path.join(BASE, 'data')
CONFIG = os.environ.get('SCRAPE_CONFIG') or os.path.join(BASE, 'config.json')

DEFAULTS = {
    'title': '招聘岗位调研汇总',
    'localCityRe': '北京',
    'localLabel': '本地',
    'localSheetName': '本地岗位',
    'rejectTarget': {},
}


def load_config():
    cfg = json.loads(json.dumps(DEFAULTS))
    try:
        with open(CONFIG, encoding='utf-8') as f:
            user = json.load(f)
        for k, v in user.items():
            if isinstance(v, dict) and isinstance(cfg.get(k), dict):
                cfg[k].update(v)
            else:
                cfg[k] = v
    except FileNotFoundError:
        pass
    except Exception as e:
        print('  ! config.json 解析失败，使用默认配置：%s' % e)
    return cfg


CFG = load_config()
TITLE = CFG['title']
LOCAL_LABEL = CFG['localLabel']
LOCAL_SHEET = CFG['localSheetName']
OUT = os.path.join(BASE, str(TITLE) + '.xlsx')

REQ_HEAD = re.compile(r'(任职要求|任职资格|岗位要求|职位要求|工作要求|任职条件|我们希望你|技能要求|资格要求|能力要求|职位要求)')
STOP_HEAD = re.compile(r'(岗位职责|工作职责|职位描述|岗位描述|工作内容|我们提供|福利待遇|薪资待遇|公司简介|加分项|其他信息|薪资范围|职位信息)')
GOODWORD = re.compile(r'(熟悉|掌握|精通|具备|了解|能够|负责|本科|硕士|专科|大专|博士|年以上|经验|能力|优先|熟练|专业|会使用|独立完成)')

# (来源标签, 文件名, 简称) —— 与各平台脚本的输出文件名一一对应
SOURCES = [
    ('猎聘',        'liepin_all.jsonl',   'lp'),
    ('前程无忧',    'job51_all.jsonl',    'j51'),
    ('智联招聘',    'zhilian_all.jsonl',  'zp'),
    ('BOSS直聘',    'boss_all.jsonl',     'boss'),
]
SRC_LABEL = {k: v for v, _, k in SOURCES}

try:
    LOCAL_CITY = re.compile(CFG['localCityRe'])
except re.error as e:
    print('  ! localCityRe 正则无效，回退为「北京」：%s' % e)
    LOCAL_CITY = re.compile('北京')

# 平台返回的「名称相似但非同一主体」的候选，人工复核后剔除（避免张冠李戴）
# 建议写在 config.json 的 rejectTarget 里；这里也可以写死，会被 config 同名键覆盖
REJECT_TARGET = {}
REJECT_TARGET.update(CFG.get('rejectTarget') or {})


def load_jsonl(name):
    """同名公司保留岗位数最多的一条；无岗位的保留一条作为兜底信息"""
    path = os.path.join(DATA, name)
    recs = {}
    if not os.path.exists(path):
        return recs
    for line in open(path, encoding='utf-8'):
        line = line.strip()
        if not line:
            continue
        try:
            d = json.loads(line)
        except Exception:
            continue
        if not isinstance(d, dict) or not d.get('n'):
            continue
        n = d['n']
        prev = recs.get(n)
        cur_j = len(d.get('jobs') or [])
        prev_j = len(prev.get('jobs') or []) if prev else -1
        if prev is None or cur_j > prev_j:
            recs[n] = d
    return recs


def clean(s):
    s = (s or '').replace('\r', '\n').replace('\xa0', ' ')
    s = re.sub(r'[ \t\u3000]+', ' ', s)
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s.strip()


def split_items(text):
    text = clean(text)
    parts = re.split(r'\n+|\s*[；;]\s*', text)
    out = []
    for p in parts:
        p = p.strip(' ·•-—*。，,、\t')
        p = re.sub(r'^[0-9０-９一二三四五六七八九十]+[、.．)）]\s*', '', p)
        p = re.sub(r'^[（(][0-9０-９一二三四五六七八九十]+[）)]\s*', '', p)
        p = re.sub(r'^[【\[][^】\]]{0,10}[】\]]\s*', '', p)
        if len(p) >= 4:
            out.append(p)
    return out


def core_req(desc, skills, jobname):
    """岗位核心要求：摘取 JD「任职要求」段落要点，50 字内"""
    d = clean(desc)
    if not d:
        sk = skills if isinstance(skills, str) else '、'.join(skills or [])
        return (sk or '（JD未明确）')[:50]
    lines = [l.strip() for l in d.split('\n') if l.strip()]

    seg, hit = [], False
    for l in lines:
        if len(l) <= 30 and REQ_HEAD.search(l):
            hit = True
            rest = REQ_HEAD.sub('', l).strip(' :：')
            if rest:
                seg.append(rest)
            continue
        if hit:
            if len(l) <= 20 and STOP_HEAD.search(l) and not re.search(r'[，,。]', l):
                break
            seg.append(l)
    if not seg:
        seg = [l for l in lines if GOODWORD.search(l)]
    items = split_items('\n'.join(seg))
    if not items:
        items = split_items('\n'.join(lines))

    out, total = [], 0
    for it in items:
        it = re.sub(r'[。；;，,]+$', '', it).strip()
        if len(it) < 4:
            continue
        if total + len(it) + 1 > 50:
            if not out:
                out.append(it[:50])
            break
        out.append(it)
        total += len(it) + 1
        if len(out) >= 3:
            break
    text = '；'.join(out)
    text = re.sub(r'^[、，,。；;\s]+', '', text)
    if len(text) > 50:
        text = text[:49].rstrip('、，,；; ') + '…'
    return text or '（JD未明确）'


def salary_split(s):
    s = (s or '').strip()
    m = re.search(r'·?\s*(\d+薪)', s)
    base = re.split(r'·', s)[0].strip()
    return base, (m.group(1) if m else '')


def unified_salary(s, src=''):
    """统一为月薪口径 K。
    各平台「万」的含义不同：猎聘的「万」= 年薪；前程无忧/智联的「万」= 月薪万（10K）。
    千/k/K = 月薪千元；元 = 月薪元。"""
    raw = (s or '').strip()
    if not raw:
        return ''
    t = re.sub(r'[·\s]*\d+\s*薪.*$', '', raw)          # 去掉 ·13薪 尾巴
    t = t.replace('～', '-').replace('~', '-').replace('—', '-').replace('至', '-')
    toks = [(float(n), u) for n, u in
            re.findall(r'(\d+(?:\.\d+)?)\s*([kK千万元])?', t) if n]
    if not toks:
        return raw
    unit = next((u for _, u in toks if u), '')
    annual = (unit == '万' and src == 'lp')

    def conv(v, u):
        u = u or unit
        if u in ('k', 'K', '千'):
            return v
        if u == '万':
            return v * 10 / 12 if annual else v * 10
        if u == '元':
            return v / 1000
        return v

    vals = [conv(v, u) for v, u in toks]
    lo, hi = vals[0], vals[-1]

    def f(x):
        return f'{x:.0f}' if (x >= 10 or abs(x - round(x)) < 0.05) else f'{x:.1f}'

    body = f'{f(lo)}-{f(hi)}K' if (len(vals) > 1 and lo != hi) else f'{f(lo)}K'
    return body + ('(年薪折算)' if annual else '')


# ---- 各平台字段适配 ----
def j_city(j, src):
    c = (j.get('city') or '').strip()
    if not c and src == 'zp':
        c = (j.get('district') or '').strip()
    return c


def j_edu(j, src):
    v = j.get('deg') or j.get('edu') or ''
    return re.sub(r'统招|全日制', '', v).strip()


def j_jd(j, src):
    return clean(j.get('jd') or j.get('desc') or '')


def j_skills(j, src):
    v = j.get('skills') or j.get('skillTags') or []
    return v if isinstance(v, str) else '、'.join([str(x) for x in v])


def j_welfare(j, src):
    w = (j.get('welfare') or '').strip()
    if not w:
        w = (j.get('welfareTag') or '').strip()
    return w or '企业未在平台填写'


def c_scale(rec, src):
    if not rec:
        return ''
    return rec.get('scale') or rec.get('size') or ''


def c_stage(rec, src):
    """公司性质/融资阶段"""
    if not rec:
        return ''
    parts = []
    for k in ('stage', 'property', 'financing'):
        v = rec.get(k)
        if v and v not in parts:
            parts.append(v)
    return ' / '.join(parts)


def c_industry(rec, src):
    return (rec or {}).get('industry') or ''


def main():
    data = {}
    for label, fname, key in SOURCES:
        data[key] = load_jsonl(fname)

    companies = json.load(open(
        os.environ.get('SCRAPE_COMPANIES') or os.path.join(BASE, 'companies.json'), encoding='utf-8'))
    order = [c['n'] for c in companies]

    rows = []        # 岗位行
    comp_rows = []   # 公司行
    miss = []
    rejected = []    # 因主体不符被剔除的命中

    for name in order:
        recs = {k: data[k].get(name) for _, _, k in SOURCES}

        # 主体校验：剔除「名称相似但非同一主体」的命中
        rej_notes = []
        pats = REJECT_TARGET.get(name)
        if pats:
            for _, _, k in SOURCES:
                r = recs[k]
                if r and r.get('target') and any(re.search(p, r['target']) for p in pats):
                    rej_notes.append(f'{SRC_LABEL[k]}命中「{r["target"]}」')
                    rejected.append([name, SRC_LABEL[k], r['target'], '疑似非同一主体，已剔除'])
                    recs[k] = None

        withjobs = {k: (r.get('jobs') or []) for k, r in recs.items() if r}

        total_jobs = sum(len(v) for v in withjobs.values())
        srcs = [SRC_LABEL[k] for _, _, k in SOURCES if recs[k] and withjobs[k]]

        if total_jobs == 0:
            detail = []
            for label, _, k in SOURCES:
                r = recs[k]
                if not r:
                    detail.append(f'{label}：无记录')
                elif r.get('notFound'):
                    cand = r.get('candidates') or []
                    detail.append(f'{label}：未命中' + (f"（候选：{'、'.join(cand[:3])}）" if cand else ''))
                else:
                    detail.append(f'{label}：0岗位')
            if rej_notes:
                detail.append('已剔除误配：' + '；'.join(rej_notes))
            miss.append({'name': name, 'detail': '；'.join(detail)})
            continue

        # 公司维度信息：按 猎聘 → 前程无忧 → 智联 → BOSS 取首个有值的
        order_keys = ['lp', 'j51', 'zp', 'boss']
        cname = next((recs[k].get('target') for k in order_keys if recs[k] and recs[k].get('target')), name)
        cscale = next((c_scale(recs[k], k) for k in order_keys if c_scale(recs[k], k)), '')
        cstage = next((c_stage(recs[k], k) for k in order_keys if c_stage(recs[k], k)), '')
        cind = next((c_industry(recs[k], k) for k in order_keys if c_industry(recs[k], k)), '')
        cnjob = max([(recs[k] or {}).get('nJobs') or 0 for k in order_keys] + [0])
        comp_rows.append([name, cname, cscale, cstage, cind, cnjob, '/'.join(srcs)] +
                         [(recs[k] or {}).get('target') or '' for k in order_keys])

        for label, _, k in SOURCES:
            rec = recs[k]
            if not rec:
                continue
            for j in (rec.get('jobs') or []):
                desc = j_jd(j, k)
                sal = j.get('salary') or ''
                base, bonus = salary_split(sal)
                rows.append([
                    name,
                    cscale or c_scale(rec, k) or '',
                    cstage or c_stage(rec, k) or '',
                    cind or c_industry(rec, k) or '',
                    j.get('jobName') or '',
                    j_city(j, k),
                    j.get('exp') or '',
                    j_edu(j, k),
                    unified_salary(sal, k) or base,
                    bonus,
                    desc,
                    core_req(desc, j_skills(j, k), j.get('jobName') or ''),
                    j_welfare(j, k),
                    label,
                ])

    if not rows:
        print('NO_ROWS')
        return

    HDR_N = 15
    for idx, r in enumerate(rows):
        if len(r) + 1 != HDR_N:
            raise SystemExit(f'列数不匹配：第 {idx} 行 {len(r)+1} 值 / 表头 {HDR_N} 列')

    wb = Workbook()
    head_fill = PatternFill('solid', fgColor='1F4E79')
    head_font = Font(color='FFFFFF', bold=True, size=10)
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    LOCAL_FILL = PatternFill('solid', fgColor='FFF2CC')

    heads = ['序号', '公司名称', '公司规模', '公司性质/融资阶段', '所属行业', '岗位名称',
             '工作城市', '经验要求', '学历要求', '薪资范围', '年薪结构',
             '岗位详细JD', '岗位核心要求（50字内）', '福利待遇', '数据来源']
    widths = [5, 17, 12, 15, 15, 28, 15, 10, 9, 13, 8, 60, 42, 26, 10]

    def style_sheet(ws, hlist, wlist, rowh=96, tab_color=None):
        for c in range(1, len(hlist) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = wlist[c - 1]
        ws.row_dimensions[1].height = 32
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(hlist) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(size=10)
            ws.cell(row=r, column=2).font = Font(size=10, bold=True)
            if rowh:
                ws.row_dimensions[r].height = rowh
        ws.freeze_panes = 'C2' if len(hlist) > 8 else 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(hlist))}{ws.max_row}'
        if tab_color:
            ws.sheet_properties.tabColor = tab_color

    # ===== Sheet1 岗位汇总（主表）=====
    ws = wb.active
    ws.title = '岗位汇总'
    ws.append(heads)
    for i, r in enumerate(rows, 1):
        ws.append([i] + r)
    style_sheet(ws, heads, widths, tab_color='1F4E79')
    # 主表中本地岗位行加底色，便于快速定位（工作城市为第 7 列）
    for r in range(2, ws.max_row + 1):
        if LOCAL_CITY.search(str(ws.cell(row=r, column=7).value or '')):
            for c in range(1, len(heads) + 1):
                ws.cell(row=r, column=c).fill = LOCAL_FILL

    # ===== Sheet2 本地岗位 =====
    ws6 = wb.create_sheet(LOCAL_SHEET)
    ws6.append(heads)
    local_rows = [r for r in rows if LOCAL_CITY.search(r[5] or '')]   # r[5]=工作城市
    for i, r in enumerate(local_rows, 1):
        ws6.append([i] + r)
    style_sheet(ws6, heads, widths, tab_color='C00000')

    # ===== Sheet3 公司概览 =====
    ws2 = wb.create_sheet('公司概览')
    h2 = ['公司名称', '招聘平台公司名', '公司规模', '公司性质/融资阶段', '所属行业', '在招岗位数(样本)',
          '数据来源', '猎聘命中名', '前程无忧命中名', '智联命中名', 'BOSS命中名']
    w2 = [18, 26, 12, 20, 18, 15, 20, 24, 24, 24, 20]
    ws2.append(h2)
    for r in comp_rows:
        ws2.append(r)
    for c in range(1, len(h2) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws2.column_dimensions[get_column_letter(c)].width = w2[c - 1]
    ws2.row_dimensions[1].height = 32
    for r in range(2, ws2.max_row + 1):
        for c in range(1, len(h2) + 1):
            ws2.cell(row=r, column=c).font = Font(size=10)
            ws2.cell(row=r, column=c).alignment = Alignment(vertical='center', wrap_text=True)
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = f'A1:{get_column_letter(len(h2))}{ws2.max_row}'
    ws2.sheet_properties.tabColor = '2E75B6'

    # ===== Sheet4 全部在招岗位 =====
    ws3 = wb.create_sheet('全部在招岗位')
    h3 = ['公司名称', '数据来源', '在招岗位数(样本)', '岗位名称', '薪资', '城市']
    ws3.append(h3)
    for name in order:
        for label, _, k in SOURCES:
            rec = data[k].get(name)
            if not rec:
                continue
            for item in rec.get('allJobs') or []:
                p = str(item).split('|')
                ws3.append([name, label, rec.get('nJobs') or '', p[0] if p else '',
                            p[1] if len(p) > 1 else '', p[3] if len(p) > 3 else (p[2] if len(p) > 2 else '')])
    for c, w in enumerate([18, 11, 15, 36, 14, 20], 1):
        cell = ws3.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws3.column_dimensions[get_column_letter(c)].width = w
    for r in range(2, ws3.max_row + 1):
        for c in range(1, 7):
            ws3.cell(row=r, column=c).font = Font(size=10)
            ws3.cell(row=r, column=c).alignment = Alignment(vertical='center')
    ws3.freeze_panes = 'A2'
    ws3.sheet_properties.tabColor = '548235'

    # ===== Sheet5 原始JD明细 =====
    ws4 = wb.create_sheet('原始JD明细')
    h4 = ['公司名称', '岗位名称', '城市', '薪资', '经验', '学历', '数据来源', '完整JD原文']
    ws4.append(h4)
    for name in order:
        for label, _, k in SOURCES:
            rec = data[k].get(name)
            if not rec:
                continue
            for j in (rec.get('jobs') or []):
                ws4.append([name, j.get('jobName'), j_city(j, k), j.get('salary'),
                            j.get('exp'), j_edu(j, k), label, j_jd(j, k)])
    for c, w in enumerate([18, 28, 16, 14, 10, 10, 11, 100], 1):
        cell = ws4.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws4.column_dimensions[get_column_letter(c)].width = w
    for r in range(2, ws4.max_row + 1):
        for c in range(1, 9):
            ws4.cell(row=r, column=c).font = Font(size=10)
            ws4.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)
        ws4.row_dimensions[r].height = 150
    ws4.freeze_panes = 'A2'
    ws4.sheet_properties.tabColor = '7F7F7F'

    # ===== Sheet6 说明与未匹配 =====
    ws5 = wb.create_sheet('说明与未匹配')
    stat = {}
    for label, _, k in SOURCES:
        stat[label] = (sum(1 for n in order if (data[k].get(n) or {}).get('jobs')),
                       sum(len((data[k].get(n) or {}).get('jobs') or []) for n in order))
    info = [
        ['数据来源', '、'.join(label for label, _, _ in SOURCES) + '（均为登录态检索）'],
        ['抓取时间', datetime.now().strftime('%Y-%m-%d %H:%M')],
        ['名单公司总数', f'{len(order)} 家'],
        ['有岗位数据公司数', f'{len(comp_rows)} 家'],
        ['岗位记录数', f'{len(rows)} 条'],
        [f'其中{LOCAL_LABEL}岗位', f'{len(local_rows)} 条（见「{LOCAL_SHEET}」表）'],
        ['', ''],
        ['各平台覆盖情况', ''],
    ]
    for label, _, _ in SOURCES:
        nco, njob = stat[label]
        info.append([f'  {label}', f'{nco} 家公司有岗位 / 共 {njob} 条岗位记录'])
    info += [
        ['', ''],
        ['字段口径说明', ''],
        ['公司规模', '招聘平台标注的员工规模区间，按 猎聘 → 前程无忧 → 智联 → BOSS 取首个有值'],
        ['公司性质/融资阶段', '如 沪深A股上市 / 已上市 / 上市公司 / B轮 等，取自招聘平台'],
        ['所属行业', '招聘平台公司页标注的行业大类（非目标产业链细分环节）'],
        ['薪资范围', '统一折算为月薪口径 K（千元/月）。各平台「万」含义不同：前程无忧、智联的「万」为月薪万（1.5-2万 = 15-20K）；猎聘的「万」为年薪（万×10÷12 折算并标注）'],
        ['年薪结构', '标注 13薪/14薪/15薪/16薪 等，空白表示按 12 薪'],
        ['岗位核心要求', '由 JD「任职要求/任职资格」段落自动提炼，严格控制在 50 字内'],
        ['福利待遇', '取自平台岗位详情页「职位福利」标签，企业未填写则标为"企业未在平台填写"'],
        ['岗位筛选逻辑', '每家公司按 研发/技术/生产/质量/工艺/职能类岗位优先级排序，取代表性岗位样本（非全量）'],
        ['匹配逻辑', '按公司名核心词严格匹配，避免同名/相似名误配；未命中时记录平台返回的候选公司名'],
        ['同公司多平台', '同一家公司在多个平台均有岗位时，各平台岗位分行并列展示，"数据来源"列区分'],
        ['', ''],
        ['未匹配到的公司', '各平台检索情况（多为该公司未在该平台发布对口岗位）'],
    ]
    for r in info:
        ws5.append(r)
    ws5.append(['公司名称', '各平台检索情况'])
    for m in miss:
        ws5.append([m['name'], m['detail']])
    if rejected:
        ws5.append(['', ''])
        ws5.append(['主体校验剔除记录', '平台返回公司名与名单公司疑似非同一主体，已剔除（避免张冠李戴）'])
        ws5.append(['公司名称', '数据来源', '平台命中公司名', '处理'])
        for r in rejected:
            ws5.append(r)
    ws5.column_dimensions['A'].width = 26
    ws5.column_dimensions['B'].width = 120
    for r in range(1, ws5.max_row + 1):
        ws5.cell(row=r, column=1).font = Font(size=10, bold=True)
        ws5.cell(row=r, column=2).font = Font(size=10)
        for c in (1, 2):
            ws5.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)
    ws5.sheet_properties.tabColor = 'BF8F00'

    wb.save(OUT)
    print('SAVED:', OUT)
    print(f'岗位行数 {len(rows)} | 有岗位公司 {len(comp_rows)} | 未匹配 {len(miss)} | {LOCAL_LABEL}岗位 {len(local_rows)}')
    for label, _, _ in SOURCES:
        nco, njob = stat[label]
        print(f'  {label}: {nco} 家 / {njob} 条')


if __name__ == '__main__':
    main()
